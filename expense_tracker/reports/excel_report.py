"""
excel_report.py
Generates a professionally formatted Excel workbook (.xlsx) using openpyxl,
including a summary sheet and a detailed transactions sheet.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.analytics_service import (
    get_summary, get_category_wise_expenses, get_transactions_dataframe
)

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F2937")
SUBTITLE_FONT = Font(size=10, italic=True, color="6B7280")
THIN_BORDER = Border(*(Side(style="thin", color="D1D5DB"),) * 4)


def export_financial_report(user_id: int, filepath: str, start_date: str = None,
                             end_date: str = None, currency: str = "Rs.") -> None:
    wb = Workbook()

    # ----- Summary Sheet -----
    ws = wb.active
    ws.title = "Summary"
    ws["B2"] = "ExpensePro - Financial Summary Report"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = f"Period: {start_date or 'All time'} to {end_date or 'Present'}"
    ws["B3"].font = SUBTITLE_FONT

    summary = get_summary(user_id, start_date, end_date)
    rows = [
        ("Total Income", f"{currency} {summary['total_income']:,.2f}"),
        ("Total Expenses", f"{currency} {summary['total_expenses']:,.2f}"),
        ("Total Savings", f"{currency} {summary['total_savings']:,.2f}"),
        ("Savings Rate", f"{summary['savings_rate']:.1f}%"),
    ]
    start_row = 5
    ws.cell(row=start_row, column=2, value="Metric").font = HEADER_FONT
    ws.cell(row=start_row, column=2).fill = HEADER_FILL
    ws.cell(row=start_row, column=3, value="Value").font = HEADER_FONT
    ws.cell(row=start_row, column=3).fill = HEADER_FILL

    for i, (label, value) in enumerate(rows, start=1):
        r = start_row + i
        ws.cell(row=r, column=2, value=label).border = THIN_BORDER
        ws.cell(row=r, column=3, value=value).border = THIN_BORDER

    # Category breakdown
    cat_start = start_row + len(rows) + 3
    ws.cell(row=cat_start, column=2, value="Category-wise Expenses").font = Font(bold=True, size=12)
    ws.cell(row=cat_start + 1, column=2, value="Category").font = HEADER_FONT
    ws.cell(row=cat_start + 1, column=2).fill = HEADER_FILL
    ws.cell(row=cat_start + 1, column=3, value="Amount").font = HEADER_FONT
    ws.cell(row=cat_start + 1, column=3).fill = HEADER_FILL

    categories = get_category_wise_expenses(user_id, start_date, end_date)
    for i, (cat, amt) in enumerate(categories.items(), start=1):
        r = cat_start + 1 + i
        ws.cell(row=r, column=2, value=cat).border = THIN_BORDER
        ws.cell(row=r, column=3, value=f"{currency} {amt:,.2f}").border = THIN_BORDER

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 22

    # ----- Transactions Sheet -----
    ws2 = wb.create_sheet("Transactions")
    df = get_transactions_dataframe(user_id, start_date, end_date)
    headers = ["ID", "Date", "Title", "Category", "Type", "Payment Method", "Amount", "Description"]
    for col, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    if not df.empty:
        for r, (_, row) in enumerate(df.iterrows(), start=2):
            values = [row["id"], row["date"], row["title"], row["category"], row["type"],
                      row["payment_method"], row["amount"], row.get("description", "")]
            for c, val in enumerate(values, start=1):
                cell = ws2.cell(row=r, column=c, value=val)
                cell.border = THIN_BORDER
                if c == 5:
                    if row["type"] == "Income":
                        cell.font = Font(color="15803D")
                    elif row["type"] == "Expense":
                        cell.font = Font(color="B91C1C")

    for col in range(1, len(headers) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 16

    wb.save(filepath)
